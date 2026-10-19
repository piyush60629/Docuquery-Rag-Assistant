
import pytest
from langchain_core.embeddings import Embeddings
from langchain_core.language_models.fake_chat_models import FakeListChatModel

from document_loader import DocumentPart
from rag_engine import (
    NOT_FOUND_MESSAGE,
    GeminiRAG,
    TextChunk,
    build_chunks,
    chunk_text,
    get_chroma_client,
)


class KeywordEmbeddings(Embeddings):
    """Deterministic offline embeddings: one dimension per keyword."""

    VOCAB = ["pipeline", "failure", "retry", "quality", "check", "incremental", "load", "watermark"]

    def _embed(self, text: str) -> list[float]:
        lowered = text.lower()
        vector = [float(lowered.count(word)) for word in self.VOCAB]
        return vector if any(vector) else [0.01] * len(self.VOCAB)

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [self._embed(text) for text in texts]

    def embed_query(self, text: str) -> list[float]:
        return self._embed(text)


@pytest.fixture()
def rag():
    llm = FakeListChatModel(responses=["The pipeline retries three times [Source 1]."])
    return GeminiRAG(embeddings=KeywordEmbeddings(), llm=llm)


CHUNKS = [
    TextChunk("runbook.txt", "On pipeline failure the job will retry three times before alerting.", None, 1),
    TextChunk("runbook.txt", "Data quality check rules validate nulls and duplicates.", None, 2),
    TextChunk("guide.pdf", "Incremental load uses a watermark column to pick new rows.", 4, 3),
]


def test_chunk_text_preserves_content_with_overlap():
    text = "Sentence one. " * 120
    chunks = chunk_text(text, chunk_size=300, overlap=50)
    assert len(chunks) > 1
    assert all(chunk.strip() for chunk in chunks)
    assert all(len(chunk) <= 300 for chunk in chunks)


def test_chunk_text_rejects_bad_settings():
    with pytest.raises(ValueError):
        chunk_text("abc", chunk_size=100)
    with pytest.raises(ValueError):
        chunk_text("abc", chunk_size=300, overlap=300)


def test_build_chunks_keeps_source_metadata():
    parts = [DocumentPart(source="runbook.txt", page=None, text="A " * 600)]
    chunks = build_chunks(parts, chunk_size=300, overlap=50)
    assert chunks
    assert all(chunk.source == "runbook.txt" for chunk in chunks)
    assert [chunk.chunk_id for chunk in chunks] == list(range(1, len(chunks) + 1))


def test_document_round_trip_keeps_page_and_handles_missing_page():
    with_page = TextChunk("guide.pdf", "text", 4, 3)
    without_page = TextChunk("notes.md", "text", None, 7)
    assert TextChunk.from_document(with_page.to_document()) == with_page
    assert "page" not in without_page.to_document().metadata
    assert TextChunk.from_document(without_page.to_document()) == without_page


def test_chroma_search_returns_best_match_first(rag):
    index = rag.build_index(CHUNKS)
    try:
        assert index.count() == 3
        results = rag.retrieve(index, "What happens on pipeline failure and retry?", top_k=2)
        assert len(results) == 2
        assert results[0].chunk.chunk_id == 1
        assert results[0].score > results[1].score
        assert 0.99 < results[0].score <= 1.0001  # cosine similarity scale

        page_hit = rag.retrieve(index, "incremental load watermark", top_k=1)[0]
        assert page_hit.chunk.label == "guide.pdf, page 4"
    finally:
        index.delete()


def test_top_k_larger_than_index_is_clamped(rag):
    index = rag.build_index(CHUNKS[:1])
    try:
        assert len(rag.retrieve(index, "pipeline", top_k=6)) == 1
    finally:
        index.delete()


def test_sessions_get_separate_collections(rag):
    first = rag.build_index(CHUNKS[:1])
    second = rag.build_index(CHUNKS[1:])
    try:
        assert first.collection_name != second.collection_name
        assert first.count() == 1 and second.count() == 2
    finally:
        first.delete()
        second.delete()
    names = {getattr(c, "name", c) for c in get_chroma_client().list_collections()}
    assert first.collection_name not in names


def test_answer_chain_uses_llm_and_handles_no_results(rag):
    index = rag.build_index(CHUNKS)
    try:
        results = rag.retrieve(index, "pipeline failure", top_k=2)
        assert rag.answer("pipeline failure?", results).startswith("The pipeline retries")
    finally:
        index.delete()
    assert rag.answer("anything", []) == NOT_FOUND_MESSAGE


def test_api_key_required_without_injected_models():
    with pytest.raises(ValueError):
        GeminiRAG(api_key="")


RUNBOOK = """Pipeline Runbook

Overview
The pipeline loads sales data every night.

Bronze Layer
Raw files land here with a batch identifier.

Failure Handling
Engineers rerun the failed batch after approval.
"""


def test_split_sections_detects_plain_text_and_markdown_headings():
    from rag_engine import split_sections

    plain = split_sections(RUNBOOK)
    assert [h for h, _ in plain] == [
        "Pipeline Runbook > Overview",
        "Pipeline Runbook > Bronze Layer",
        "Pipeline Runbook > Failure Handling",
    ]
    assert plain[1][1] == "Raw files land here with a batch identifier."

    markdown = split_sections("# Guide\nIntro text.\n\n## Setup\nInstall it.\n\n### Windows\nUse PowerShell.\n\n## Usage\nRun it.")
    assert [h for h, _ in markdown] == ["Guide", "Guide > Setup", "Guide > Setup > Windows", "Guide > Usage"]


def test_split_sections_does_not_treat_sentences_or_bullets_as_headings():
    from rag_engine import split_sections

    text = "Intro line that ends with a period.\n\n- Bullet item\nmore text\n\nShort sentence.\nNext line."
    assert split_sections(text) == [("", text)]


def test_section_strategy_keeps_topics_apart_and_prefixes_headings():
    parts = [DocumentPart(source="runbook.txt", text=RUNBOOK)]
    fixed = build_chunks(parts, chunk_size=300, overlap=50, strategy="fixed")
    sections = build_chunks(parts, chunk_size=300, overlap=50, strategy="sections")
    assert len(fixed) == 1  # the whole runbook fits one fixed-size chunk, mixing three topics
    assert len(sections) == 3
    assert sections[1].text.startswith("Pipeline Runbook > Bronze Layer\n")
    assert sections[1].section == "Pipeline Runbook > Bronze Layer"
    assert "rerun" not in sections[1].text
    assert TextChunk.from_document(sections[1].to_document()) == sections[1]


def test_long_sections_are_split_within_the_size_limit():
    body = "Each sentence describes one more validation rule. " * 40
    parts = [DocumentPart(source="rules.md", text=f"## Validation Rules\n{body}")]
    chunks = build_chunks(parts, chunk_size=400, overlap=50, strategy="sections")
    assert len(chunks) > 1
    assert all(len(c.text) <= 400 for c in chunks)
    assert all(c.text.startswith("Validation Rules\n") for c in chunks)


def test_unknown_strategy_is_rejected():
    with pytest.raises(ValueError):
        build_chunks([DocumentPart(source="a.txt", text="x")], strategy="magic")


def test_docx_headings_become_markdown_sections():
    from io import BytesIO

    from docx import Document as WordDocument

    from document_loader import extract_docx
    from rag_engine import split_sections

    word = WordDocument()
    word.add_heading("Onboarding Guide", level=0)
    word.add_paragraph("Welcome to the team.")
    word.add_heading("Access", level=1)
    word.add_paragraph("Raise a request in the portal.")
    buffer = BytesIO()
    word.save(buffer)

    parts = extract_docx(buffer.getvalue(), "guide.docx")
    assert parts[0].text.startswith("# Onboarding Guide\n\nWelcome")
    assert [h for h, _ in split_sections(parts[0].text)] == ["Onboarding Guide", "Onboarding Guide > Access"]


def test_llm_provider_resolution():
    from rag_engine import resolve_llm_provider

    assert resolve_llm_provider(None, groq_api_key="") == "gemini"
    assert resolve_llm_provider(None, groq_api_key="gsk_x") == "groq"
    assert resolve_llm_provider("Gemini", groq_api_key="gsk_x") == "gemini"
    with pytest.raises(ValueError):
        resolve_llm_provider("openai")


def test_groq_and_gemini_chat_models_are_built_with_the_same_settings():
    from langchain_google_genai import ChatGoogleGenerativeAI
    from langchain_groq import ChatGroq

    from rag_engine import DEFAULT_GROQ_MODEL, make_chat_model

    groq_rag = GeminiRAG(api_key="gemini-key", llm_provider="groq", llm_api_key="groq-key", max_retries=0, request_timeout=60)
    assert isinstance(groq_rag.llm, ChatGroq)
    assert groq_rag.llm.model_name == DEFAULT_GROQ_MODEL == groq_rag.generation_model
    assert groq_rag.llm.max_retries == 0 and groq_rag.llm.request_timeout == 60

    assert isinstance(GeminiRAG(api_key="gemini-key").llm, ChatGoogleGenerativeAI)
    judge = make_chat_model("groq", "groq-key", "openai/gpt-oss-120b", temperature=0)
    assert judge.reasoning_effort == "low"


def test_groq_provider_needs_its_own_key():
    with pytest.raises(ValueError, match="Groq"):
        GeminiRAG(api_key="gemini-key", llm_provider="groq", llm_api_key="")


def test_reasoning_models_get_headroom_and_hidden_thinking():
    from rag_engine import REASONING_TOKEN_HEADROOM, make_chat_model

    answer = make_chat_model("groq", "k", "openai/gpt-oss-120b", max_tokens=700)
    judge = make_chat_model("groq", "k", "qwen/qwen3.8-27b", max_tokens=300)
    assert answer.max_tokens == 700 + REASONING_TOKEN_HEADROOM and answer.reasoning_effort == "low"
    assert judge.reasoning_format == "hidden" and judge.max_tokens == 300 + REASONING_TOKEN_HEADROOM
